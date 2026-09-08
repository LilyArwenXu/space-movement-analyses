"""Shared page chrome, workbook-faithful categories and information chart controls."""
from pathlib import Path
import re,html
from build_data import ROOT

def category_page():
    import openpyxl
    path=ROOT/'category-photos/街道空间要素分类表_最新.xlsx'
    ws=openpyxl.load_workbook(path,data_only=True).active
    starts={};covered=set()
    for m in ws.merged_cells.ranges:
        starts[(m.min_row,m.min_col)]=(m.max_row-m.min_row+1,m.max_col-m.min_col+1)
        covered.update((r,c) for r in range(m.min_row,m.max_row+1) for c in range(m.min_col,m.max_col+1) if (r,c)!=(m.min_row,m.min_col))
    codes=['A1a','A2a','A3','A4a','B1a','B1b','B1c','B1d','B2a','B2b','B2c','B2d','B3a-i','B3a-ii','B3b-i','B3b-ii','B3b-iii','B3cc','C1','C2','C3','C4','D1','D2','D3','D4','D5','D6','D7','D8']
    photos=list((ROOT/'category-photos').glob('*.png'));rows=[]
    for r in range(1,ws.max_row+1):
        cells=[]
        for c in range(1,ws.max_column+1):
            if (r,c) in covered:continue
            h,w=starts.get((r,c),(1,1));value=str(ws.cell(r,c).value or '');images=[]
            eligible=(c>=2 if r<=4 or r>=19 else c>=4)
            if eligible:
                for rr in range(r,r+h):
                    code=codes[rr-1];photo=next((p for p in photos if p.stem==code or p.stem.endswith('_'+code)),None)
                    if photo and photo.name not in images:images.append(photo.name)
            popup='<span class="classification-popup">'+''.join(f'<img loading="lazy" src="category-photos/{html.escape(p)}" alt="{html.escape(value)}">' for p in images)+'</span>' if images else ''
            attrs=' tabindex="0"' if images else ''
            cells.append(f'<td rowspan="{h}" colspan="{w}"{attrs}>{html.escape(value)}{popup}</td>')
        rows.append('<tr>'+''.join(cells)+'</tr>')
    return '<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>类型划分</title><link rel="stylesheet" href="assets/css/inclusive.css"></head><body><header class="topbar"><a href="index.html">←</a><h1>类型划分</h1></header><main class="classification"><table aria-label="街道空间要素分类表"><colgroup>'+''.join('<col>' for _ in range(7))+'</colgroup><tbody>'+''.join(rows)+'</tbody></table></main></body></html>'

def update_renderer(js):
    js=re.sub(r"const FONT='[^']*'", "const FONT='SimSun, Songti SC, serif'", js)
    js=js.replace("font:'16px Microsoft YaHei'", "font:'16px '+FONT")
    js=js.replace("font:'12px SimSun'","font:'12px '+FONT").replace("font:'11px SimSun'","font:'11px '+FONT").replace("font:'14px SimSun'","font:'14px '+FONT").replace("font:'16px SimSun'","font:'16px '+FONT")
    for start,end in [('function informationScatter(', 'function mixingInformation('),('function informationExplorer(', 'function scatter(')]:
        a=js.index(start);b=js.index(end,a);part=js[a:b]
        part=part.replace('{\n','{\n const points=rankedInformationPoints(kind);\n',1).replace('D.points','points')
        if start.startswith('function informationExplorer'):
            part=part.replace("const controls=el('div','metric-controls',host)","const controls=el('div','metric-controls axis-controls',host)")
            part=part.replace("main.textContent=label;", "main.textContent='';main.title=label;main.setAttribute('aria-label',label);main.style.background=PALETTE[j%PALETTE.length];")
            part=part.replace("controls.hidden=!mode;", "controls.hidden=!mode;positionAxisControls(controls,c,count);c.reflow=()=>positionAxisControls(controls,c,count);")
            part=part.replace('bottom:50}', 'bottom:82}')
        part=part.replace('按总表顺序','按综合评分从低到高').replace('总表点位顺序','综合评分排序').replace('每列为一个点位','每列为一个点位（按综合评分升序，缺失评分置后）')
        js=js[:a]+part+js[b:]
    js=js.replace("if(PAGE==='memory')mainTabs(['在地记忆的假设','活力度信息表'],i=>{if(i)informationExplorer(W,'vitality');else", "if(PAGE==='memory')mainTabs(['活力度信息表','在地记忆的假设'],i=>{if(!i)informationExplorer(W,'vitality');else")
    js=js.replace("charts.push(c);return c;", "charts.push(c);enableChartDownload(c,n);return c;")
    return js+'\n'+(ROOT/'aerial/assets/js/interface-tools.js').read_text(encoding='utf-8')

def update_pages():
    for path in (ROOT/'aerial').glob('*.html'):
        s=path.read_text(encoding='utf-8')
        if path.name=='index.html':
            s=s.replace('<body','<body data-cover="true"',1) if 'data-cover=' not in s else s
        elif path.name!='visualizations.html':
            title=re.search(r'<h1>(.*?)</h1>',s)
            if title:s=re.sub(r'<header class="topbar">.*?</header>',f'<header class="topbar unified-topbar"><a class="brand" href="index.html">← INCLUSIVE VITALITY</a><h1>{title[1]}</h1></header>',s,flags=re.S)
        if 'interface.css' not in s:s=s.replace('</head>','<link rel="stylesheet" href="assets/css/interface.css"></head>')
        path.write_text(s,encoding='utf-8')
    for directory in ['spacemovement','aerial/spacemovement']:
        for path in (ROOT/directory).glob('*.html'):
            s=path.read_text(encoding='utf-8')
            if 'data-page=' not in s:continue
            s=re.sub(r'<a class="back"[^>]*>.*?</a>', '<a class="brand" href="../visualizations.html">← INCLUSIVE VITALITY</a>',s)
            s=s.replace('class="topbar"','class="topbar unified-topbar"')
            if 'interface.css' not in s:s=s.replace('</head>','<link rel="stylesheet" href="../assets/css/interface.css"></head>')
            path.write_text(s,encoding='utf-8')
